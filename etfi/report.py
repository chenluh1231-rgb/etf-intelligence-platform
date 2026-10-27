"""Excel 报告导出（原生图表，便于在 PPT / 微信等渠道二次使用）。

偏好：用 Excel 原生图表而非图片，方便复用与再编辑。
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import EXPORT_DIR

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=13, bold=True, color="1F4E79")


def _sheet(wb, title, headers, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        ws.cell(1, c).fill = HEAD_FILL
        ws.cell(1, c).font = HEAD_FONT
        ws.cell(1, c).alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, 1):
        w = (widths or {}).get(h, max(10, len(str(h)) * 2 + 4))
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


def export(user_id: str, profile: dict, match: dict | None,
           brief: dict | None) -> str:
    wb = Workbook()
    wb.remove(wb.active)
    rev = (profile or {}).get("revealed") or {}

    # ---- 1. 组合与画像 ----
    pos_rows = []
    for h in (brief or {}).get("portfolio_pulse", {}).get("holdings", []):
        pos_rows.append([h["code"], h["name"], h["weight"], h["pct"], h["contribution"]])
    if not pos_rows:
        for k, v in (rev.get("asset_allocation") or {}).items():
            pos_rows.append(["", k, v, None, None])
    ws1 = _sheet(wb, "组合与画像", ["代码", "名称", "权重", "当日涨跌%", "贡献%"], pos_rows,
                 {"名称": 28})
    r = len(pos_rows) + 3
    ws1.cell(r, 1, "资产配置").font = TITLE_FONT
    r += 1
    ws1.cell(r, 1, "资产类别")
    ws1.cell(r, 2, "占比")
    rr = r + 1
    for i, (k, v) in enumerate((rev.get("asset_allocation") or {}).items()):
        ws1.cell(rr + i, 1, k)
        ws1.cell(rr + i, 2, v)
    if rev.get("asset_allocation"):
        pie = PieChart()
        pie.title = "组合资产配置"
        pie.add_data(Reference(ws1, min_col=2, min_row=r, max_row=rr + len(rev["asset_allocation"]) - 1),
                     titles_from_data=True)
        pie.set_categories(Reference(ws1, min_col=1, min_row=rr,
                                     max_row=rr + len(rev["asset_allocation"]) - 1))
        pie.height, pie.width = 8, 12
        ws1.add_chart(pie, "H2")

    # ---- 2. 底层穿透 ----
    sec = rev.get("sector_exposure") or []
    ws2 = _sheet(wb, "底层穿透", ["行业", "穿透权重"], [[s["sector"], s["weight"]] for s in sec],
                 {"行业": 20})
    if sec:
        bar = BarChart()
        bar.type = "bar"
        bar.title = "穿透后行业暴露（前十大持仓口径）"
        bar.add_data(Reference(ws2, min_col=2, min_row=1, max_row=len(sec) + 1), titles_from_data=True)
        bar.set_categories(Reference(ws2, min_col=1, min_row=2, max_row=len(sec) + 1))
        bar.height, bar.width = 10, 16
        ws2.add_chart(bar, "E2")
    r2 = len(sec) + 3
    ws2.cell(r2, 1, "穿透后个股前 15 大").font = TITLE_FONT
    for i, s in enumerate((rev.get("top_stocks") or [])[:15], 1):
        ws2.cell(r2 + i, 1, s["code"])
        ws2.cell(r2 + i, 2, s["name"])
        ws2.cell(r2 + i, 3, s["weight"])

    # ---- 3. 匹配结果 ----
    rows3 = []
    for res_type, label in (("best_match", "Best Match"), ("best_addition", "Best Portfolio Addition")):
        for r3 in (match or {}).get(res_type) or []:
            c = r3.get("components") or {}
            rows3.append([label, r3.get("rank"), r3.get("code"), r3.get("name"),
                          r3.get("score"), r3.get("addition_score"),
                          c.get("goal"), c.get("preference"), c.get("complement"),
                          c.get("cost"), c.get("risk"),
                          r3.get("overlap_norm"), r3.get("exposure_gain"),
                          r3.get("corr_with_portfolio"), r3.get("marginal_vol_change_pct"),
                          r3.get("fee"), r3.get("scale_yi"),
                          " | ".join(r3.get("reasons") or []),
                          " | ".join(r3.get("mismatches") or [])])
    ws3 = _sheet(wb, "匹配结果",
                 ["结果类型", "排名", "代码", "名称", "匹配分", "增量分", "目标", "偏好",
                  "互补", "成本", "风险", "与组合重叠", "新增暴露", "相关性",
                  "边际波动变化%", "费率%", "规模(亿)", "匹配理由", "不匹配因素"],
                 rows3, {"名称": 26, "匹配理由": 60, "不匹配因素": 60})
    if rows3:
        bm = [x for x in rows3 if x[0] == "Best Match"][:10]
        if bm:
            base = 1
            start = 2
            ws3.cell(1, 21, "图表辅助列")
            for i, x in enumerate(bm):
                ws3.cell(1 + i + 1, 21, x[4])
            bar2 = BarChart()
            bar2.title = "Best Match 匹配得分"
            bar2.add_data(Reference(ws3, min_col=21, min_row=1, max_row=len(bm) + 1),
                          titles_from_data=True)
            bar2.set_categories(Reference(ws3, min_col=4, min_row=start, max_row=len(bm) + 1))
            bar2.height, bar2.width = 9, 16
            ws3.add_chart(bar2, "W2")

    # ---- 4. 组合内重叠 ----
    pairs = rev.get("overlap_pairs") or []
    ws4 = _sheet(wb, "持仓重叠", ["ETF A", "代码A", "ETF B", "代码B", "归一化重叠", "共同持仓占净值"],
                 [[p.get("a_name", p["a"]), p["a"], p.get("b_name", p["b"]), p["b"],
                   p["overlap"], p.get("raw")] for p in pairs],
                 {"ETF A": 24, "ETF B": 24})

    # ---- 5. 每日简报 ----
    if brief:
        lines = [brief.get("headline", "")]
        for s in brief.get("signals") or []:
            lines.append("· " + s)
        for x in (brief.get("what_matters") or {}).get("etf_moves", [])[:10]:
            lines.append(f"· {x['name']}（{x['code']}）当日 {x['pct']:+.2f}%，"
                         f"权重 {x['weight']:.1%}{('，' + x['why']) if x.get('why') else ''}")
        for x in (brief.get("etf_radar") or {}).get("new_products", [])[:6]:
            lines.append(f"· 新产品 {x['name']}（{x['code']}）上市 {x['age_days']} 天，"
                         f"规模 {x.get('scale_yi')} 亿元")
        ws5 = _sheet(wb, "每日简报", [f"每日 ETF 情报 · {brief.get('trade_date')}"],
                     [[x] for x in lines], {"A": 90})
        ws5.cell(len(lines) + 3, 1, brief.get("disclaimer", "")).font = Font(italic=True, size=9)

    path = EXPORT_DIR / f"ETF情报_{user_id}_{(brief or {}).get('trade_date', 'latest')}.xlsx"
    wb.save(path)
    return str(path)
